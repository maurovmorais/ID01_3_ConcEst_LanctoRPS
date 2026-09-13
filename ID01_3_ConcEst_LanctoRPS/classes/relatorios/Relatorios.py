# Imports dos módulos internos do projeto
# Carrega o InitAllSettingsSettings Precisa ser o primeiro a ser carregado
from ID01_3_ConcEst_LanctoRPS.classes.framework.InitAllSettings import InitAllSettings
from ID01_3_ConcEst_LanctoRPS.classes.utils.ExecutionControl import ExecutionControl
from ID01_3_ConcEst_LanctoRPS.classes.dados_execucao.DadosExecucao import DadosExecucao

# Imports dos pacotes externos
from openpyxl import load_workbook
from datetime import datetime
import shutil, os
import re
import pandas as pd


class Relatorios:
    """
    Classe responsável por manipular relatórios, incluindo linhas e preenchendo a partir de templates
    
    Parâmetros:

    Retorna:

    """
    _config = InitAllSettings.config
    _guid_execucao = InitAllSettings.guid_execucao
    caminho_bd_analit_sint = InitAllSettings.sqlite_caminho_bd_analit_sint
    


    @classmethod
    def preencher_analitico(cls) -> str:
        """
        Realiza o preenchimento do excel do relatorio analitico.
        
        Parâmetros:
        
        
        Retorna:
            - str: Caminho do arquivo preenchido com os dados do relatorio analitico.
        
        """
        cls.path_relatorio_analitico:str = os.path.join(cls._config["CaminhoPastaRelatorios"] , "Relatorio_Analitico_" + re.sub(string=str(cls._config["NomeProcesso"]),pattern=r'[\@\$%&\\\/\:\*\?\"\'<>\|~`#\^\+=\{\}\[\];\!]',repl='') + "___datetime__" + ".xlsx")
        cls.path_relatorio_analitico = cls.path_relatorio_analitico.replace("__datetime__",datetime.now().strftime("%d%m%Y%H%M%S"))

        #Copiando apenas se não existir
        if(not os.path.exists(cls.path_relatorio_analitico)):
            shutil.copy(src=InitAllSettings.caminho_template_excel_analitico, dst=cls.path_relatorio_analitico)
        
            #Colocando nome do processo no relatório
            wbk_analitico = load_workbook(cls.path_relatorio_analitico)
            wsht_analitico = wbk_analitico.active
            wsht_analitico["D4"] = cls._config["NomeProcesso"]

            wbk_analitico.save(cls.path_relatorio_analitico)
            wbk_analitico.close()
        
        # Executando o comando select
        dados_analitico = DadosExecucao.get_dados_analitico()['Dados']

        # Nome aba arquivo excel que será realizada a leitura
        sheet_name = '2. Analítico'

        # Carregar o arquivo Excel
        wb_excel_analitico = load_workbook(cls.path_relatorio_analitico)

        # Selecionar a aba Analitico
        ws_aba_analitico = wb_excel_analitico[sheet_name]

        # Variaveis auxiliares para encontrar linha vazia no excel                
        index_newline:int = None
        index_aux = 5

        # Encontrando linha vazia
        while(index_newline is None):
            if(ws_aba_analitico["A" + index_aux.__str__()].value is None):
                index_newline = index_aux
            else:
                index_aux += 1

        # Percorre linhas e colunas do excel para preencher com os dados analiticos
        for r_idx, row in enumerate(dados_analitico, 1):
            for c_idx, value in enumerate(row, 1):
                ws_aba_analitico.cell(row=r_idx+index_newline-1, column=c_idx, value=value)

        # Salva as modificações
        wb_excel_analitico.save(cls.path_relatorio_analitico)

        return cls.path_relatorio_analitico 

    @classmethod    
    def preencher_sintetico(cls) -> str:
        """
        Realiza o preenchimento do excel do relatorio sintetico.
        
        Parâmetros:
        
        
        Retorna:
            - str: Caminho do arquivo preenchido com os dados do relatorio sintetico.
        
        """
        cls.path_relatorio_sintetico:str = os.path.join(cls._config["CaminhoPastaRelatorios"] , "Relatorio_Sintetico_" + re.sub(string=str(cls._config["NomeProcesso"]),pattern=r'[\@\$%&\\\/\:\*\?\"\'<>\|~`#\^\+=\{\}\[\];\!]',repl='') + "___datetime__" + ".xlsx")
        cls.path_relatorio_sintetico = cls.path_relatorio_sintetico.replace("__datetime__",datetime.now().strftime("%d%m%Y%H%M%S"))

        if(not os.path.exists(cls.path_relatorio_sintetico)):
            shutil.copy(src=InitAllSettings.caminho_template_excel_sintetico, dst=cls.path_relatorio_sintetico)
        
            #Colocando nome do processo no relatório
            wbk_sintetico = load_workbook(cls.path_relatorio_sintetico)
            wsht_sintetico = wbk_sintetico.active
            wsht_sintetico["C4"] = cls._config["NomeProcesso"]

            wbk_sintetico.save(cls.path_relatorio_sintetico)
            wbk_sintetico.close()
      
        # Executando o comando select
        dados_sintetico = DadosExecucao.get_dados_sintetico()['Dados']

        # Nome aba arquivo excel que será realizada a leitura
        sheet_name = '1. Sintético'

        # Carregar o arquivo Excel
        wb_excel_sintetico = load_workbook(cls.path_relatorio_sintetico)

        # Selecionar a aba Sintetico
        ws_aba_sintetico= wb_excel_sintetico[sheet_name]

        # Variaveis auxiliares para encontrar linha vazia no excel                
        index_newline:int = None
        index_aux = 5

        # Encontrando linha vazia
        while(index_newline is None):
            if(ws_aba_sintetico["A" + index_aux.__str__()].value is None):
                index_newline = index_aux
            else:
                index_aux += 1

        # Percorre linhas e colunas do excel para preencher com os dados sinteticos
        for r_idx, row in enumerate(dados_sintetico, 1):
            for c_idx, value in enumerate(row, 1):
                ws_aba_sintetico.cell(row=r_idx+index_newline-1, column=c_idx, value=value)

        # Salva as modificações
        wb_excel_sintetico.save(cls.path_relatorio_sintetico)

        return cls.path_relatorio_sintetico     
    
    @classmethod    
    def preencher_csv_sintetico_banco_raas(cls) -> str:
        """
        Realiza o preenchimento do csv que é o relatorio sintetico para o banco raas.
        
        Parâmetros:
        
        
        Retorna:
            - str: Caminho do arquivo preenchido com os dados do relatorio sintetico para o banco raas.
        
        """
        path_relatorio_sintetico_raas:str = os.path.join(cls._config["CaminhoPastaRelatorios"] , "Relatorio_Sintetico_RAAS_" + re.sub(string=str(cls._config["NomeProcesso"]),pattern=r'[\@\$%&\\\/\:\*\?\"\'<>\|~`#\^\+=\{\}\[\];\!]',repl='') + "___datetime__" + ".csv")


        # Renomeando arquivo para incluir data hora minuto e segundo no nome do arquivo
        caminho_arquivo_csv_com_data = path_relatorio_sintetico_raas.replace("__datetime__",datetime.now().strftime("%d%m%Y%H%M%S"))
        
        col_dados_sintetico = DadosExecucao.get_dados_raas_sintetico()
        linha_capturada_fila = col_dados_sintetico['Dados']
        name_columns:list = col_dados_sintetico['Colunas']

        dados_capturados_fila = None

        if(linha_capturada_fila is not None):
            dados_capturados_fila:dict = dict(zip(name_columns, linha_capturada_fila))
        
        if ExecutionControl.task is not None:
            status_job = 'PARADO' if ExecutionControl.is_interrupted() else 'SUCESSO'
        else:
            status_job = 'SUCESSO'

        if cls._config['job_id'] is None:
            tipo_execucao = 'MANUAL'
        elif cls._config['job_id'] == 0:
            tipo_execucao = 'MANUAL_VSCODE'
        else:
            tipo_execucao = 'AUTOMATICO'

        # Adicionando dados que não existem na tabela de execucao
        dados_capturados_fila['Ferramenta'] = 'Python/Selenium'
        dados_capturados_fila['Cliente'] = cls._config["NomeCliente"]
        dados_capturados_fila['tenant'] = cls._config.get("Ambiente", "")
        dados_capturados_fila['pasta'] = None
        dados_capturados_fila['descri_processo'] = cls._config["DescricaoProcesso"]
        dados_capturados_fila['Serv_Captcha'] = 'S' if (InitAllSettings.usa_captcha) else 'N'
        dados_capturados_fila['Qtd_captcha'] = InitAllSettings.qtde_captcha
        dados_capturados_fila['Serv_Ocr'] = 'S' if (InitAllSettings.usa_ocr) else 'N'
        dados_capturados_fila['Qtd_Ocr'] = InitAllSettings.qtde_ocr
        dados_capturados_fila['Serv_Api'] = 'S' if (InitAllSettings.usa_api) else 'N'
        dados_capturados_fila['Qtd_Api'] = InitAllSettings.qtde_api
        dados_capturados_fila['job_id'] = ExecutionControl.job_id
        dados_capturados_fila['status_job'] = status_job
        dados_capturados_fila['tipo_execucao'] = tipo_execucao

        # Rename columns
        name_columns_in_raas_database = {'nome_maquina':'Nome_maquina',
                                             'resolucao_tela':'Resolucao',
                                             'inicio_exec':'inicio',
                                             'fim_exec':'fim',
                                             'qtd_itens_fila':'total_itens',
                                             'qtd_itens_sucesso':'total_itens_sucesso',
                                             'qtd_itens_business':'total_itens_negocios',
                                             'qtd_itens_app':'total_itens_aplicacao'}
        

        df_dados_sinteticos = pd.DataFrame(dados_capturados_fila, index=[0])
        df_dados_sinteticos.rename(columns=name_columns_in_raas_database, inplace=True)


        column_order = ['Ferramenta','Cliente','tenant','pasta',
                               'nome_processo','descri_processo','Nome_maquina','Resolucao',
                               'Serv_Captcha','Qtd_captcha','Serv_Ocr','Qtd_Ocr',
                               'Serv_Api','Qtd_Api','inicio','fim',
                               'total_itens','total_itens_sucesso','total_itens_negocios','total_itens_aplicacao','job_id','status_job','tipo_execucao']
        df_dados_sinteticos = df_dados_sinteticos.reindex(columns=column_order)


        df_dados_sinteticos.to_csv(path_or_buf=caminho_arquivo_csv_com_data,sep=';',encoding='utf-8',index=False)


        return caminho_arquivo_csv_com_data

    @classmethod    
    def preencher_csv_analitico_banco_raas(cls) -> str:
        """
        Realiza o preenchimento do csv que é o relatorio analitico para o banco raas.
        
        Parâmetros:
        
        
        Retorna:
            - str: Caminho do arquivo preenchido com os dados do relatorio analitico para o banco raas.
        
        """
        path_relatorio_sintetico_raas:str = os.path.join(cls._config["CaminhoPastaRelatorios"] , "Relatorio_Analitico_RAAS_" + re.sub(string=str(cls._config["NomeProcesso"]),pattern=r'[\@\$%&\\\/\:\*\?\"\'<>\|~`#\^\+=\{\}\[\];\!]',repl='') + "___datetime__" + ".csv")


        # Renomeando arquivo para incluir data hora minuto e segundo no nome do arquivo
        caminho_arquivo_csv_com_data = path_relatorio_sintetico_raas.replace("__datetime__",datetime.now().strftime("%d%m%Y%H%M%S"))
        

        linha_capturada_fila:list = DadosExecucao.get_dados_raas_analitico()['Dados']


        column_order = ['nome_fila','referencia','item_fila','inicio',
                               'fim','status_item_fila','tipo_excecao','descr_excecao']
        df_dados_analiticos = pd.DataFrame(linha_capturada_fila,columns=column_order)


        df_dados_analiticos.to_csv(path_or_buf=caminho_arquivo_csv_com_data,sep=';',encoding='utf-8',index=False)


        return caminho_arquivo_csv_com_data


